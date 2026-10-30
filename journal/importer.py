"""Broker CSV / XLSX import with format auto-detection.

Supported out of the box
------------------------
* Zerodha Console **Tradebook** export (one row per fill, BUY and SELL mixed)
* Any broker "fills" CSV: date + symbol + side + qty + price (columns can be
  named differently or ordered differently)
* Round-trip exports: entry date/price + exit date/price (Zerodha P&L, custom)
* This app's own CSV export (round-trip with notes/tags)

The parser never assumes a fixed column order. It sniffs the delimiter, finds
the header row even below title/preamble lines, maps columns by alias, then
either feeds fills to the FIFO engine or creates closed trades directly.
"""
from __future__ import annotations

import csv
import io
import json
import re
from typing import Any

from .db import get_settings
from .trades import apply_fill, create_trade
from .utils import norm_side, norm_symbol, parse_date, parse_num

# ---------------------------------------------------------------------------
# column aliases
# ---------------------------------------------------------------------------

ALIASES: dict[str, list[str]] = {
    "symbol": [
        "symbol", "tradingsymbol", "trading_symbol", "trading symbol", "scrip", "scripname",
        "scrip name", "stock", "stock name", "instrument", "instrument name", "security",
        "security name", "company", "name of the instrument", "underlying", "code",
    ],
    "side": [
        "trade_type", "trade type", "transaction_type", "transaction type", "buy_sell",
        "buy/sell", "side", "action", "order_type", "type", "buy or sell", "transaction",
    ],
    "quantity": [
        "quantity", "qty", "qty.", "traded_qty", "traded quantity", "filled_quantity",
        "executed_quantity", "net_quantity", "net qty", "shares", "units", "no. of shares",
        "lot size", "lots",
    ],
    "price": [
        "price", "trade_price", "trade price", "execution_price", "execution price",
        "average_price", "avg price", "avg. price", "avg_price", "rate", "trade rate",
        "buy_average", "buy average", "buy_avg", "buy price", "buy_price", "purchase price",
        "entry_price", "entry price", "entry", "avg buy", "buy rate",
    ],
    "date": [
        "trade_date", "trade date", "execution_date", "execution date", "order_execution_time",
        "order execution time", "entry_date", "entry date", "buy_date", "buy date",
        "purchase_date", "date", "trade time", "time", "datetime", "executed_at",
    ],
    "exit_date": [
        "exit_date", "exit date", "sell_date", "sell date", "sale_date", "sale date",
        "date_of_sale", "closing_date", "close_date", "sold date",
    ],
    "exit_price": [
        "exit_price", "exit price", "exit", "sell_price", "sell price", "sell_average",
        "sell average", "sell_avg", "sell rate", "sale_price", "sale price", "avg sell",
        "closing_price", "close price",
    ],
    "fees": [
        "fees", "fee", "charges", "total_charges", "total charges", "brokerage",
        "brokerage_charges", "taxes", "total tax & charges", "total taxes and charges",
        "stt", "transaction charges", "gst",
    ],
    "order_id": ["order_id", "order id", "orderid", "order_no", "order no", "order_number"],
    "trade_id": ["trade_id", "trade id", "tradeid", "trade_no", "trade no", "execution_id", "fill_id", "trade number"],
    "exchange": ["exchange", "exch", "market"],
    "segment": ["segment", "product", "product_type", "series"],
    "notes": ["notes", "note", "remarks", "comment", "comments", "reason", "setup"],
    "strategy": ["strategy", "system", "plan", "setup name"],
    "tags": ["tags", "tag", "label", "labels"],
    "status": ["status", "result", "outcome", "trade_status"],
    "pnl": ["pnl", "p&l", "profit", "profit/loss", "profit_loss", "realised_pnl", "realized_pnl",
            "net_pnl", "net p&l", "net profit", "gross pnl", "realised p&l", "realized p&l"],
    "stop_loss": ["stop_loss", "stop loss", "sl", "stoploss", "stop"],
    "target": ["target", "tp", "take profit", "target price"],
    "rating": ["rating", "stars", "score"],
}

_NORM_HEADER = re.compile(r"[^a-z0-9&/ ]+")

# headers that hint a round-trip (closed trade) export rather than raw fills
_EXIT_HEADERS = {"exit_price", "exit_date"}


def _norm_header(cell: str) -> str:
    s = str(cell or "").strip().lower().replace("_", " ")
    s = _NORM_HEADER.sub(" ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


HEADER_LOOKUP: dict[str, str] = {}
for field, names in ALIASES.items():
    for n in names:
        HEADER_LOOKUP.setdefault(_norm_header(n), field)
        HEADER_LOOKUP.setdefault(_norm_header(n).replace(" ", "_"), field)


# ---------------------------------------------------------------------------
# reading
# ---------------------------------------------------------------------------

def decode_bytes(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "utf-16", "utf-16-le", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            if text.count("\ufffd") <= max(2, len(text) // 1000):
                return text
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def detect_delimiter(text: str) -> str:
    sample_lines = [ln for ln in text.splitlines()[:25] if ln.strip()][:15]
    sample = "\n".join(sample_lines)
    best, best_score = ",", -1
    for delim in (",", ";", "\t", "|"):
        counts = [ln.count(delim) for ln in sample_lines]
        if not counts:
            continue
        freq = sum(1 for c in counts if c > 0)
        score = freq * 10 + (sum(counts) / max(1, len(counts)))
        if score > best_score:
            best, best_score = delim, score
    return best if best_score > 0 else ","


def _read_rows(text: str, delimiter: str) -> list[list[str]]:
    try:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter, skipinitialspace=True)
        return [row for row in reader]
    except csv.Error:
        return []


def _score_header(row: list[str]) -> int:
    score = 0
    for cell in row:
        if _norm_header(cell) in HEADER_LOOKUP:
            score += 1
    return score


def find_header(rows: list[list[str]]) -> tuple[int, int]:
    best_idx, best_score = 0, -1
    for i, row in enumerate(rows[:30]):
        if not row or all(not str(c).strip() for c in row):
            continue
        score = _score_header(row)
        if score > best_score:
            best_idx, best_score = i, score
        if score >= 3:
            break
    return best_idx, max(best_score, 0)


def map_columns(header: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header):
        key = _norm_header(cell)
        field = HEADER_LOOKUP.get(key) or HEADER_LOOKUP.get(key.replace(" ", "_"))
        if field and field not in mapping:
            mapping[field] = idx
        elif not field:
            # tolerate "Buy Price (₹)" style suffixes
            for known, f in HEADER_LOOKUP.items():
                if len(known) >= 4 and key.startswith(known) and f not in mapping:
                    mapping[f] = idx
                    break
    return mapping


def _detect_dayfirst(values: list[str]) -> bool:
    dayfirst_votes = 0
    monthfirst_votes = 0
    for v in values[:200]:
        m = re.match(r"\s*(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})", str(v or ""))
        if not m:
            continue
        a, b = int(m.group(1)), int(m.group(2))
        if a > 12 >= b:
            dayfirst_votes += 1
        elif b > 12 >= a:
            monthfirst_votes += 1
    if monthfirst_votes > dayfirst_votes:
        return False
    return True


def _cell(row: list[str], mapping: dict[str, int], field: str) -> str | None:
    idx = mapping.get(field)
    if idx is None or idx >= len(row):
        return None
    val = row[idx]
    return str(val).strip() if val is not None else None


def _repair_row(row: list[str], header_len: int, numeric_fields: list[int]) -> list[str]:
    """Some exports contain unquoted 1,234.56 numbers which break CSV columns."""
    if len(row) <= header_len:
        return row
    out = list(row)
    idx = 0
    while len(out) > header_len and idx < len(out) - 1:
        nxt = out[idx + 1]
        if (
            idx in numeric_fields
            and re.fullmatch(r"-?\d{1,3}", out[idx].strip() or "")
            and re.fullmatch(r"\d{3}(\.\d+)?", nxt.strip() or "")
        ):
            out[idx : idx + 2] = [f"{out[idx]},{nxt}"]
            continue
        idx += 1
    return out


# ---------------------------------------------------------------------------
# parse
# ---------------------------------------------------------------------------

def parse_csv(data: bytes, filename: str = "") -> dict[str, Any]:
    text = decode_bytes(data)
    delimiter = detect_delimiter(text)
    rows = _read_rows(text, delimiter)
    rows = [r for r in rows if r and any(str(c).strip() for c in r)]
    if not rows:
        return {
            "ok": False,
            "error": "The file looks empty.",
            "filename": filename,
            "delimiter": delimiter,
            "rows": [],
        }

    header_idx, score = find_header(rows)
    header = [str(c).strip() for c in rows[header_idx]]
    mapping = map_columns(header)
    body = rows[header_idx + 1 :]

    if score < 2 or "symbol" not in mapping:
        return {
            "ok": False,
            "error": (
                "Could not find a symbol column. The file may not be a trade export, or the "
                "header row is unusual. You can still add trades manually."
            ),
            "filename": filename,
            "delimiter": delimiter,
            "header": header,
            "mapping": mapping,
            "rows": [],
        }

    date_col = mapping.get("date")
    date_values = []
    if date_col is not None:
        date_values = [r[date_col] for r in body if date_col < len(r)]
    dayfirst = _detect_dayfirst(date_values)

    numeric_fields = [mapping[f] for f in ("price", "quantity", "exit_price", "fees", "pnl") if f in mapping]
    header_len = len(header)
    fills: list[dict[str, Any]] = []
    roundtrips: list[dict[str, Any]] = []
    problems: list[dict[str, Any]] = []
    skipped = 0

    # Decide once whether this is a round-trip export or a fill export.
    is_roundtrip = ("exit_price" in mapping or "exit_date" in mapping) and "price" in mapping
    is_roundtrip = is_roundtrip and "side" not in mapping
    has_side = "side" in mapping

    for i, row in enumerate(body):
        if len(row) > header_len:
            row = _repair_row(row, header_len, numeric_fields)
        first = " ".join(str(c).strip().lower() for c in row[:2])
        if first.startswith("total") or first.startswith("grand total") or first.startswith("disclaimer"):
            skipped += 1
            continue
        symbol = norm_symbol(_cell(row, mapping, "symbol"))
        if not symbol or symbol in ("TOTAL", "GRANDTOTAL", "NAN", "NONE"):
            skipped += 1
            continue
        qty = parse_num(_cell(row, mapping, "quantity"))
        price = parse_num(_cell(row, mapping, "price"))
        date = parse_date(_cell(row, mapping, "date"), dayfirst=dayfirst)
        exit_price = parse_num(_cell(row, mapping, "exit_price"))
        exit_date = parse_date(_cell(row, mapping, "exit_date"), dayfirst=dayfirst)
        pnl = parse_num(_cell(row, mapping, "pnl"))
        fees = parse_num(_cell(row, mapping, "fees")) or 0.0
        side = norm_side(_cell(row, mapping, "side"))
        status = norm_side(_cell(row, mapping, "status"))

        if qty is None or qty == 0:
            problems.append({"row": i + header_idx + 2, "reason": "missing quantity", "raw": row[:8]})
            continue

        if is_roundtrip:
            if price is None or not date:
                problems.append({"row": i + header_idx + 2, "reason": "missing entry price/date", "raw": row[:8]})
                continue
            if exit_price is None and pnl is not None and qty:
                # back-compute exit price from P&L when the export omits it
                exit_price = price + (pnl / qty)
            rt = {
                "symbol": symbol,
                "quantity": abs(qty),
                "entry_date": date,
                "entry_price": price,
                "exit_date": exit_date or date,
                "exit_price": exit_price,
                "fees": fees,
                "notes": _cell(row, mapping, "notes") or "",
                "strategy": _cell(row, mapping, "strategy") or "",
                "tags": _cell(row, mapping, "tags") or "",
                "exchange": (_cell(row, mapping, "exchange") or "NSE").upper(),
                "external_key": _cell(row, mapping, "trade_id") or _cell(row, mapping, "order_id"),
                "source": "CSV",
            }
            if status == "OPEN":
                rt["exit_price"] = None
                rt["exit_date"] = None
            roundtrips.append(rt)
            continue

        if price is None or not date:
            problems.append({"row": i + header_idx + 2, "reason": "missing price/date", "raw": row[:8]})
            continue
        if not side:
            # infer from a P&L/status column for single-sided exports
            side = "SELL" if status == "LOSS" else "BUY"
        direction = "SELL" if "SELL" in str(side) else "BUY"
        fills.append(
            {
                "symbol": symbol,
                "side": direction,
                "quantity": abs(qty),
                "price": price,
                "exec_date": date,
                "fees": fees,
                "exchange": (_cell(row, mapping, "exchange") or "NSE").upper(),
                "order_id": _cell(row, mapping, "order_id"),
                "trade_ref": _cell(row, mapping, "trade_id"),
                "notes": _cell(row, mapping, "notes") or "",
                "strategy": _cell(row, mapping, "strategy") or "",
                "_row": i + header_idx + 2,
            }
        )

    kind = "roundtrip" if roundtrips else "fills"
    normalized = roundtrips if roundtrips else fills
    return {
        "ok": True,
        "filename": filename,
        "delimiter": delimiter,
        "header_row": header_idx + 1,
        "header": header,
        "mapping": {k: header[v] for k, v in mapping.items()},
        "kind": kind,
        "has_side": has_side,
        "dayfirst": dayfirst,
        "row_count": len(normalized),
        "skipped": skipped,
        "problems": problems[:50],
        "preview": normalized[:12],
        "rows": normalized,
        "columns_detected": sorted(mapping.keys()),
    }


# ---------------------------------------------------------------------------
# import
# ---------------------------------------------------------------------------

def import_parsed(conn, parsed: dict[str, Any], settings: dict[str, Any] | None = None) -> dict[str, Any]:
    settings = settings or get_settings(conn)
    summary: dict[str, Any] = {
        "filename": parsed.get("filename"),
        "kind": parsed.get("kind"),
        "total_rows": len(parsed.get("rows") or []),
        "opened": 0,
        "closed": 0,
        "created_trades": 0,
        "duplicates": 0,
        "review": 0,
        "ignored": 0,
        "errors": [],
        "details": [],
    }
    rows = parsed.get("rows") or []
    if parsed.get("kind") == "roundtrip":
        for rt in rows:
            payload = {k: v for k, v in rt.items() if k != "source"}
            res = create_trade(conn, payload, source="CSV", settings=settings)
            if res.get("created"):
                summary["created_trades"] += 1
            elif res.get("duplicate"):
                summary["duplicates"] += 1
                summary["details"].append({"row": rt["symbol"], "action": "duplicate", "message": res.get("message")})
            else:
                summary["review"] += 1
                summary["details"].append({"row": rt["symbol"], "action": "review", "message": res.get("message")})
        summary["imported"] = summary["created_trades"]
    else:
        for fill in rows:
            try:
                res = apply_fill(conn, fill, source="CSV", settings=settings)
            except Exception as exc:  # keep going, report the row
                summary["errors"].append({"row": fill.get("_row"), "error": str(exc)})
                summary["ignored"] += 1
                continue
            action = res.get("action")
            if action == "opened":
                summary["opened"] += 1
                summary["created_trades"] += 1
            elif action == "closed":
                summary["closed"] += 1
            elif action == "partial_review":
                summary["closed"] += 1
                summary["review"] += 1
            elif action == "duplicate":
                summary["duplicates"] += 1
            elif action == "review":
                summary["review"] += 1
            else:
                summary["ignored"] += 1
            if action in ("duplicate", "review", "partial_review", "ignored"):
                summary["details"].append(
                    {
                        "row": fill.get("_row"),
                        "symbol": fill.get("symbol"),
                        "side": fill.get("side"),
                        "action": action,
                        "message": res.get("reason") or res.get("message") or "",
                    }
                )
        summary["imported"] = summary["created_trades"] + summary["closed"]
    return summary


def parse_xlsx(data: bytes) -> dict[str, Any]:
    """Optional XLSX support (needs openpyxl). Converts first sheet to CSV text."""
    try:
        import openpyxl  # type: ignore
    except ImportError:
        return {
            "ok": False,
            "error": "XLSX needs the openpyxl package (pip install openpyxl). Please export CSV instead.",
            "rows": [],
        }
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        buf = io.StringIO()
        writer = csv.writer(buf)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(["" if c is None else c for c in row])
        parsed = parse_csv(buf.getvalue().encode("utf-8"), "workbook.xlsx")
        return parsed
    except Exception as exc:  # pragma: no cover
        return {"ok": False, "error": f"Could not read XLSX: {exc}", "rows": []}
